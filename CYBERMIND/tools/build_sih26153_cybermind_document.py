from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "SIH26153_CYBERMIND_Idea_Abstract_Architecture.docx"
ARCH = ROOT / "releases" / "CYBERMIND_ARCHITECTURE_2026-09-29.png"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=100, bottom=85, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    r = p.add_run(text)
    r.font.size = Pt(8.6)
    return p


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(8.8)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.line_spacing = 1.0

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(4)

    h1 = styles["Heading 1"]
    h1.font.name = "Aptos Display"
    h1.font.size = Pt(14)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.space_before = Pt(7)
    h1.paragraph_format.space_after = Pt(3)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Aptos Display"
    h2.font.size = Pt(11)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.space_before = Pt(5)
    h2.paragraph_format.space_after = Pt(2)
    h2.paragraph_format.keep_with_next = True


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("CYBERMIND Project Brief and Architecture")

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(5)
    meta_run = meta.add_run("Smart India Hackathon 2026  |  Problem Statement SIH26153  |  Predictive Cyber Defence")
    meta_run.bold = True
    meta_run.font.size = Pt(9.0)
    meta_run.font.color.rgb = RGBColor(45, 62, 80)

    doc.add_heading("Idea Description", level=1)
    add_body(doc, "CYBERMIND is an offline, analyst-centred predictive cyber-defence platform that moves beyond the conventional detect-and-alert model. It represents a network as a time-ordered sequence of host communication graphs, learns how that state changes, and forecasts the risk and coarse attack stage of upcoming graph windows. This gives defenders an evidence-linked view of what may happen next while an intrusion is still developing.")
    add_body(doc, "An analyst imports a PCAP, PCAPNG, CAP, or network-flow CSV. CYBERMIND parses the input locally, preserves file provenance, converts records into a canonical flow schema, and builds overlapping temporal graph windows in which hosts are nodes and communications are edges. A pinned graph-temporal model then combines a GATv2 graph encoder, temporal Transformer, learned latent dynamics, and multi-head prediction layers to generate a short-horizon risk trajectory, coarse stage estimates, and uncertainty indicators.")
    add_body(doc, "The forecast is placed inside a practical investigation workflow. The Command Centre links network topology, observed flows, timelines, and forecasts. Parallel Futures evaluates permitted in-memory changes such as host isolation, port blocking, or rate limiting and compares their predicted future risk with the no-action baseline. These outputs are decision-support simulations; CYBERMIND does not silently enforce controls on a production network. A bounded offline assistant answers read-only questions about the loaded case, while local validation records actual sandbox HTTP responses and SHA-256 hashes beside a frozen forecast.")
    add_body(doc, "For secure collaboration, the Reports module creates an incident brief and encrypts it locally using AES-256-GCM with a fresh random key and nonce. The encrypted JSON does not contain the secret key. A recipient imports the file into CYBERMIND, enters the separately shared key, and can view or download the readable report. The complete system runs locally as a Windows desktop application or from prebuilt offline Docker images, keeping the core analysis pipeline independent of cloud APIs.")

    doc.add_heading("Why the Solution Stands Out", level=2)
    add_bullet(doc, "Forecasts future graph-window risk instead of classifying isolated flows only.")
    add_bullet(doc, "Connects prediction to counterfactual defence comparisons and analyst approval.")
    add_bullet(doc, "Keeps observations, rule-based flow triage, model forecasts, and validation evidence explicitly separate.")
    add_bullet(doc, "Operates offline and supports authenticated, encrypted incident-report exchange.")

    doc.add_heading("Project Abstract", level=1)
    add_body(doc, "Modern intrusion-detection systems mainly describe traffic that has already occurred, leaving defenders little support for anticipating attack progression or comparing possible responses. CYBERMIND addresses SIH26153 through an offline graph-temporal world model for short-horizon cyber-risk forecasting. The platform transforms packet captures or flow records into rolling host-communication graphs, encodes spatial relationships with GATv2, models temporal evolution with a Transformer, and rolls a learned latent dynamics model forward to estimate future infiltration risk and coarse attack stage. The forecasts are integrated with an analyst workspace that provides network visualisation, rule-based flow triage, uncertainty cues, counterfactual intervention comparisons, bounded case questions, sandbox-response validation, and AES-256-GCM incident reporting. The selected checkpoint achieved pooled held-out binary F1 of 0.9849, precision of 0.9970, recall of 0.9730, and false-positive rate of 0.0056 on the documented capture-grouped CIC-IDS2018 protocol, outperforming the feature-matched logistic baseline on that split. CYBERMIND therefore demonstrates a practical shift from retrospective alerting toward evidence-linked prediction, simulation, and informed defensive planning, while clearly treating stage estimates and simulated interventions as decision aids rather than verified ground truth or autonomous enforcement.")

    doc.add_section(WD_SECTION.NEW_PAGE)
    section2 = doc.sections[-1]
    section2.top_margin = Inches(0.48)
    section2.bottom_margin = Inches(0.45)
    section2.left_margin = Inches(0.55)
    section2.right_margin = Inches(0.55)

    heading = doc.add_paragraph(style="Title")
    heading.add_run("System Architecture")
    intro = doc.add_paragraph("The architecture keeps ingestion, learned forecasting, analyst decision support, validation, and protected reporting within one local application boundary.")
    intro.paragraph_format.space_after = Pt(4)

    image_p = doc.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.space_after = Pt(4)
    image_p.add_run().add_picture(str(ARCH), width=Inches(7.15))

    doc.add_heading("Architecture Layers", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.15), Inches(2.2), Inches(3.85)]
    headers = ["Layer", "Core Components", "Responsibility and Output"]
    for idx, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(8.7)
    set_repeat_table_header(table.rows[0])

    rows = [
        ("1  Ingestion", "Local parser, provenance, normaliser, graph builder", "Converts PCAP or CSV into source-tagged canonical flows and overlapping host-communication graph windows. Rule-based flow triage remains separate from model predictions."),
        ("2  Model Core", "GATv2, temporal Transformer, latent dynamics, risk and stage heads", "Encodes recent graph history and produces multi-step graph-window risk, coarse stage, predictive dispersion, model identity, and optional attribution evidence."),
        ("3  Analyst Layer", "Command Centre, Threat Forecast, Parallel Futures, assistant", "Links evidence to forecasts, compares in-memory defensive alternatives, ranks reversible suggestions, and answers bounded questions about the loaded case."),
        ("4  Trust and Delivery", "Local validator, AES-256-GCM reports, Windows and Docker runtimes", "Stores bounded response evidence beside a frozen forecast, encrypts incident briefs locally, supports in-app decryption, and keeps core processing offline."),
    ]
    for ridx, row_data in enumerate(rows):
        cells = table.add_row().cells
        for cidx, (cell, text) in enumerate(zip(cells, row_data)):
            cell.width = widths[cidx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ridx % 2:
                set_cell_shading(cell, "EDF4F8")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if cidx else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(8.1)
            if cidx == 0:
                run.bold = True

    doc.add_heading("End to End Data Flow", level=2)
    flow = doc.add_paragraph()
    flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    flow.paragraph_format.space_after = Pt(3)
    r = flow.add_run("PCAP or CSV  ->  canonical flows  ->  temporal graphs  ->  best.pt forecast  ->  analyst simulation and validation  ->  encrypted incident brief")
    r.bold = True
    r.font.size = Pt(9)

    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(0)
    note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    lead = note.add_run("Operational boundaries. ")
    lead.bold = True
    note.add_run("The model forecasts aggregate graph windows, not verified malicious labels for individual flows. Parallel Futures compares model outcomes on modified graph copies and does not prove causal impact. Local validation records target responses but does not retrain the checkpoint. Consequential defensive actions remain under analyst control.")

    for sec in doc.sections:
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.space_before = Pt(2)
        footer_run = footer.add_run("SIH26153  |  CYBERMIND  |  Project Brief and Architecture")
        footer_run.font.size = Pt(7.5)
        footer_run.font.color.rgb = RGBColor(90, 90, 90)

    doc.core_properties.title = "CYBERMIND Project Brief and Architecture"
    doc.core_properties.subject = "SIH26153 idea description abstract and two-page architecture"
    doc.core_properties.author = "CYBERMIND Team"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
